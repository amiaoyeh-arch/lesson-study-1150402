# -*- coding: utf-8 -*-
import os
import sys
import json
import argparse
import subprocess
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_html_deck(slides_data, output_html_path, title_text="學習共同體課例研究分享簡報"):
    slides_json_str = json.dumps(slides_data, ensure_ascii=False, indent=2)
    
    html_content = f'''<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title_text}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Zen+Maru+Gothic:wght@400;500;700;900&family=Noto+Sans+TC:wght@400;500;700&display=swap" rel="stylesheet">
    <style>
        :root {{
            --bg-color: #F7F5F0;
            --text-color: #2D3748;
            --primary-accent: #2E5A44;
            --secondary-accent: #4A7C59;
            --highlight: #D97724;
            --card-bg: #FFFFFF;
            --border-color: #E2DDD5;
            --subtext-color: #5A6A7E;
            --font-family: 'Zen Maru Gothic', 'Noto Sans TC', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
        }}

        [data-theme="dark"] {{
            --bg-color: #1A202C;
            --text-color: #E2E8F0;
            --primary-accent: #68D391;
            --secondary-accent: #48BB78;
            --highlight: #F6AD55;
            --card-bg: #2D3748;
            --border-color: #4A5568;
            --subtext-color: #A0AEC0;
        }}

        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
        }}

        body {{
            font-family: var(--font-family);
            background-color: var(--bg-color);
            color: var(--text-color);
            display: flex;
            justify-content: center;
            align-items: center;
            min-height: 100vh;
            overflow: hidden;
            transition: background-color 0.3s, color 0.3s;
        }}

        .deck-container {{
            width: 100vw;
            height: 100vh;
            max-width: 177.78vh;
            max-height: 56.25vw;
            aspect-ratio: 16 / 9;
            position: relative;
            background: var(--bg-color);
            box-shadow: 0 10px 30px rgba(0,0,0,0.08);
            display: flex;
            flex-direction: column;
            overflow: hidden;
        }}

        .deck-header {{
            height: clamp(36px, 6vh, 56px);
            padding: 0 clamp(16px, 2.5vw, 36px);
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-bottom: 1px solid var(--border-color);
            background: rgba(255, 255, 255, 0.4);
            backdrop-filter: blur(8px);
            z-index: 10;
        }}

        .deck-title {{
            font-size: clamp(0.75rem, 1.1vw, 0.95rem);
            font-weight: 700;
            color: var(--primary-accent);
            letter-spacing: 0.05em;
        }}

        .deck-controls {{
            display: flex;
            gap: 10px;
            align-items: center;
        }}

        .btn {{
            background: transparent;
            border: 1px solid var(--border-color);
            color: var(--text-color);
            padding: 4px 12px;
            border-radius: 6px;
            font-size: clamp(0.7rem, 1vw, 0.85rem);
            font-family: inherit;
            cursor: pointer;
            transition: all 0.2s;
        }}

        .btn:hover {{
            background: var(--primary-accent);
            color: white;
            border-color: var(--primary-accent);
        }}

        .slide-stage {{
            flex: 1;
            position: relative;
            overflow: hidden;
            padding: clamp(12px, 2vh, 24px) clamp(16px, 2.5vw, 36px);
        }}

        .slide {{
            position: absolute;
            top: clamp(12px, 2vh, 24px);
            bottom: clamp(12px, 2vh, 24px);
            left: clamp(16px, 2.5vw, 36px);
            right: clamp(16px, 2.5vw, 36px);
            display: grid;
            grid-template-columns: 1.15fr 0.85fr;
            gap: clamp(16px, 2.2vw, 32px);
            opacity: 0;
            visibility: hidden;
            transition: opacity 0.3s cubic-bezier(0.4, 0, 0.2, 1), transform 0.3s cubic-bezier(0.4, 0, 0.2, 1);
            transform: translateX(15px);
        }}

        .slide.active {{
            opacity: 1;
            visibility: visible;
            transform: translateX(0);
        }}

        .slide-left {{
            display: flex;
            flex-direction: column;
            justify-content: flex-start;
            height: 100%;
            overflow: hidden;
        }}

        .badge {{
            display: inline-block;
            align-self: flex-start;
            padding: 3px 10px;
            background: rgba(46, 90, 68, 0.1);
            color: var(--primary-accent);
            border-radius: 4px;
            font-size: clamp(0.65rem, 0.9vw, 0.8rem);
            font-weight: 700;
            margin-bottom: clamp(4px, 0.8vh, 8px);
            border-left: 3px solid var(--primary-accent);
            flex-shrink: 0;
        }}

        .slide-title {{
            font-size: clamp(1.15rem, 2vw, 1.7rem);
            font-weight: 900;
            color: var(--primary-accent);
            line-height: 1.25;
            margin-bottom: clamp(2px, 0.5vh, 6px);
            flex-shrink: 0;
        }}

        .slide-subtitle {{
            font-size: clamp(0.75rem, 1.1vw, 0.95rem);
            color: var(--subtext-color);
            margin-bottom: clamp(8px, 1.2vh, 14px);
            font-weight: 500;
            line-height: 1.35;
            flex-shrink: 0;
        }}

        .slide-content-card {{
            background: var(--card-bg);
            border-radius: 12px;
            padding: clamp(12px, 1.8vh, 22px) clamp(14px, 1.8vw, 24px);
            border: 1px solid var(--border-color);
            box-shadow: 0 4px 15px rgba(0,0,0,0.03);
            flex: 1;
            display: flex;
            flex-direction: column;
            justify-content: center;
            overflow-y: auto;
            max-height: 100%;
        }}

        .slide-content-card::-webkit-scrollbar {{
            width: 4px;
        }}
        .slide-content-card::-webkit-scrollbar-thumb {{
            background: var(--border-color);
            border-radius: 4px;
        }}

        .bullets-list {{
            display: flex;
            flex-direction: column;
            list-style: none;
            width: 100%;
            margin: auto 0;
        }}

        .bullet-item {{
            position: relative;
            padding-left: 1.2em;
            line-height: 1.45;
            color: var(--text-color);
            font-size: clamp(0.75rem, 1.05vw, 0.98rem);
        }}

        .bullet-item::before {{
            content: "•";
            position: absolute;
            left: 0.1em;
            color: var(--highlight);
            font-weight: bold;
            font-size: 1.1em;
        }}

        .bullet-item.sub-item {{
            padding-left: 2.2em;
            font-size: 0.92em;
            color: var(--subtext-color);
        }}

        .bullet-item.sub-item::before {{
            content: "–";
            left: 1.2em;
            color: var(--secondary-accent);
        }}

        .slide-right {{
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            height: 100%;
            position: relative;
        }}

        .image-card {{
            width: 100%;
            height: 100%;
            background: var(--card-bg);
            border-radius: 12px;
            overflow: hidden;
            border: 1px solid var(--border-color);
            box-shadow: 0 4px 15px rgba(0,0,0,0.04);
            display: flex;
            align-items: center;
            justify-content: center;
            position: relative;
        }}

        .image-card img {{
            width: 100%;
            height: 100%;
            object-fit: cover;
            transition: transform 0.3s ease;
        }}

        .image-card:hover img {{
            transform: scale(1.02);
        }}

        .deck-footer {{
            height: clamp(32px, 5vh, 48px);
            padding: 0 clamp(16px, 2.5vw, 36px);
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-top: 1px solid var(--border-color);
            background: rgba(255, 255, 255, 0.4);
            backdrop-filter: blur(8px);
            font-size: clamp(0.65rem, 0.9vw, 0.8rem);
            color: var(--subtext-color);
            z-index: 10;
        }}

        .progress-bar-container {{
            flex: 1;
            height: 4px;
            background: var(--border-color);
            border-radius: 2px;
            margin: 0 20px;
            overflow: hidden;
        }}

        .progress-bar-fill {{
            height: 100%;
            background: var(--primary-accent);
            width: 0%;
            transition: width 0.3s ease;
        }}

        .nav-arrow {{
            background: var(--card-bg);
            border: 1px solid var(--border-color);
            color: var(--text-color);
            width: 28px;
            height: 28px;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            transition: all 0.2s;
        }}

        .nav-arrow:hover {{
            background: var(--primary-accent);
            color: white;
        }}

        .notes-drawer {{
            position: absolute;
            bottom: clamp(32px, 5vh, 48px);
            left: 0;
            right: 0;
            background: rgba(255, 255, 255, 0.95);
            backdrop-filter: blur(10px);
            border-top: 2px solid var(--primary-accent);
            padding: 16px clamp(16px, 2.5vw, 36px);
            transform: translateY(100%);
            transition: transform 0.3s cubic-bezier(0.4, 0, 0.2, 1);
            z-index: 20;
            box-shadow: 0 -5px 20px rgba(0,0,0,0.1);
            max-height: 40vh;
            overflow-y: auto;
        }}

        [data-theme="dark"] .notes-drawer {{
            background: rgba(45, 55, 72, 0.95);
        }}

        .notes-drawer.open {{
            transform: translateY(0);
        }}

        .notes-title {{
            font-size: 0.85rem;
            font-weight: 700;
            color: var(--primary-accent);
            margin-bottom: 6px;
        }}

        .notes-content {{
            font-size: 0.85rem;
            line-height: 1.6;
            color: var(--text-color);
        }}
    </style>
</head>
<body>
    <div class="deck-container" id="deckContainer">
        <header class="deck-header">
            <div class="deck-title" id="deckHeaderTitle">{title_text}</div>
            <div class="deck-controls">
                <button class="btn" id="btnNotes">講稿 (N)</button>
                <button class="btn" id="btnTheme">主題 (T)</button>
                <button class="btn" id="btnFullscreen">全螢幕 (F)</button>
            </div>
        </header>

        <main class="slide-stage" id="slideStage"></main>

        <div class="notes-drawer" id="notesDrawer">
            <div class="notes-title">講稿要點</div>
            <div class="notes-content" id="notesContent"></div>
        </div>

        <footer class="deck-footer">
            <div id="slideCounter">1 / 1</div>
            <div class="progress-bar-container">
                <div class="progress-bar-fill" id="progressFill"></div>
            </div>
            <div style="display: flex; gap: 8px;">
                <button class="nav-arrow" id="btnPrev" title="上一頁 (←)">‹</button>
                <button class="nav-arrow" id="btnNext" title="下一頁 (→)">›</button>
            </div>
        </footer>
    </div>

    <script>
        const slidesData = {slides_json_str};
        let currentSlideIndex = 0;
        const totalSlides = slidesData.length;

        const slideStage = document.getElementById('slideStage');
        const slideCounter = document.getElementById('slideCounter');
        const progressFill = document.getElementById('progressFill');
        const notesDrawer = document.getElementById('notesDrawer');
        const notesContent = document.getElementById('notesContent');
        const btnNotes = document.getElementById('btnNotes');
        const btnTheme = document.getElementById('btnTheme');
        const btnFullscreen = document.getElementById('btnFullscreen');
        const btnPrev = document.getElementById('btnPrev');
        const btnNext = document.getElementById('btnNext');

        function cleanBulletText(text) {{
            return text.replace(/^[\\s•·–\\-*\\d\\.\\)]+\\s*/, '');
        }}

        function renderSlides() {{
            slideStage.innerHTML = '';
            slidesData.forEach((data, idx) => {{
                const slide = document.createElement('div');
                slide.className = `slide ${{idx === 0 ? 'active' : ''}}`;
                slide.id = `slide-${{idx}}`;

                let bulletsHtml = '';
                data.bullets.forEach(b => {{
                    const isSub = b.trim().startsWith('•') || b.trim().startsWith('-') || b.trim().startsWith('–') || (b.startsWith('  ') || b.startsWith('    '));
                    const cleanText = cleanBulletText(b);
                    bulletsHtml += `<li class="bullet-item ${{isSub ? 'sub-item' : ''}}">${{cleanText}}</li>`;
                }});

                slide.innerHTML = `
                    <div class="slide-left">
                        <div class="badge">${{data.badge || '課例研究'}}</div>
                        <h1 class="slide-title">${{data.title}}</h1>
                        <div class="slide-subtitle">${{data.subtitle || ''}}</div>
                        <div class="slide-content-card">
                            <ul class="bullets-list">
                                ${{bulletsHtml}}
                            </ul>
                        </div>
                    </div>
                    <div class="slide-right">
                        <div class="image-card">
                            <img src="${{data.image}}" alt="${{data.title}}" loading="lazy">
                        </div>
                    </div>
                `;
                slideStage.appendChild(slide);
            }});
        }}

        function autoFitSlideText(slideElem) {{
            if (!slideElem) return;
            const card = slideElem.querySelector('.slide-content-card');
            const list = slideElem.querySelector('.bullets-list');
            if (!card || !list) return;

            const cardHeight = card.clientHeight;
            if (cardHeight <= 0) return;

            let fontSize = 16;
            let gap = 12;
            list.style.fontSize = `${{fontSize}}px`;
            list.style.gap = `${{gap}}px`;

            let listHeight = list.scrollHeight;
            let targetHeight = cardHeight * 0.88;

            while (listHeight > targetHeight && fontSize > 11) {{
                fontSize -= 0.5;
                gap = Math.max(4, gap - 0.8);
                list.style.fontSize = `${{fontSize}}px`;
                list.style.gap = `${{gap}}px`;
                listHeight = list.scrollHeight;
            }}

            if (listHeight < cardHeight * 0.65 && fontSize < 19) {{
                while (listHeight < cardHeight * 0.82 && fontSize < 20) {{
                    fontSize += 0.5;
                    gap = Math.min(16, gap + 0.8);
                    list.style.fontSize = `${{fontSize}}px`;
                    list.style.gap = `${{gap}}px`;
                    listHeight = list.scrollHeight;
                }}
            }}
        }}

        function updateSlide(newIndex) {{
            if (newIndex < 0 || newIndex >= totalSlides) return;
            const slides = document.querySelectorAll('.slide');
            slides[currentSlideIndex].classList.remove('active');
            currentSlideIndex = newIndex;
            slides[currentSlideIndex].classList.add('active');

            slideCounter.innerText = `${{currentSlideIndex + 1}} / ${{totalSlides}}`;
            progressFill.style.width = `${{((currentSlideIndex + 1) / totalSlides) * 100}}%`;
            notesContent.innerText = slidesData[currentSlideIndex].notes || '無講稿內容';

            requestAnimationFrame(() => {{
                autoFitSlideText(slides[currentSlideIndex]);
            }});
        }}

        btnPrev.addEventListener('click', () => updateSlide(currentSlideIndex - 1));
        btnNext.addEventListener('click', () => updateSlide(currentSlideIndex + 1));

        btnNotes.addEventListener('click', () => {{
            notesDrawer.classList.toggle('open');
        }});

        btnTheme.addEventListener('click', () => {{
            const currentTheme = document.documentElement.getAttribute('data-theme');
            const newTheme = currentTheme === 'dark' ? 'light' : 'dark';
            document.documentElement.setAttribute('data-theme', newTheme);
        }});

        btnFullscreen.addEventListener('click', () => {{
            if (!document.fullscreenElement) {{
                document.documentElement.requestFullscreen().catch(err => console.log(err));
            }} else {{
                document.exitFullscreen();
            }}
        }});

        window.addEventListener('resize', () => {{
            const activeSlide = document.querySelector('.slide.active');
            autoFitSlideText(activeSlide);
        }});

        document.addEventListener('keydown', (e) => {{
            if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') {{
                updateSlide(currentSlideIndex + 1);
            }} else if (e.key === 'ArrowLeft' || e.key === 'PageUp') {{
                updateSlide(currentSlideIndex - 1);
            }} else if (e.key.toLowerCase() === 'n') {{
                notesDrawer.classList.toggle('open');
            }} else if (e.key.toLowerCase() === 't') {{
                btnTheme.click();
            }} else if (e.key.toLowerCase() === 'f') {{
                btnFullscreen.click();
            }}
        }});

        renderSlides();
        updateSlide(0);
    </script>
</body>
</html>'''

    with open(output_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"[SUCCESS] HTML Deck generated at: {output_html_path}")


def build_pptx_deck(slides_data, output_pptx_path, base_assets_dir=""):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(247, 245, 240)
    PRIMARY_COLOR = RGBColor(46, 90, 68)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 221, 213)
    TEXT_COLOR = RGBColor(45, 55, 72)
    SUBTEXT_COLOR = RGBColor(90, 106, 126)
    ACCENT_BADGE = RGBColor(230, 240, 235)
    BULLET_ORANGE = RGBColor(217, 119, 36)

    def set_shape_flat_color(shape, fill_color, border_color=None):
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()

    for s_idx, data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        set_shape_flat_color(bg, BG_COLOR)

        # Header bar
        header_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12.133), Inches(0.4))
        tf_h = header_box.text_frame
        tf_h.word_wrap = True
        p_h = tf_h.paragraphs[0]
        p_h.text = "學習共同體課例研究成果報告"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.name = "微軟正黑體"
        p_h.font.color.rgb = PRIMARY_COLOR

        # Slide Badge
        badge_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.72), Inches(3.0), Inches(0.35))
        set_shape_flat_color(badge_shape, ACCENT_BADGE, PRIMARY_COLOR)
        tf_b = badge_shape.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = data.get("badge", "課例研究")
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.name = "微軟正黑體"
        p_b.font.color.rgb = PRIMARY_COLOR
        p_b.alignment = PP_ALIGN.CENTER

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.12), Inches(6.8), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = data.get("title", "")
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.name = "微軟正黑體"
        p_t.font.color.rgb = PRIMARY_COLOR

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.95), Inches(6.8), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_top = tf_sub.margin_bottom = tf_sub.margin_left = tf_sub.margin_right = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data.get("subtitle", "")
        p_sub.font.size = Pt(12)
        p_sub.font.name = "微軟正黑體"
        p_sub.font.color.rgb = SUBTEXT_COLOR

        # Content Card
        card_w, card_h = Inches(6.7), Inches(4.35)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.45), card_w, card_h)
        set_shape_flat_color(card, CARD_BG, BORDER_COLOR)

        bullets_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(6.3), Inches(4.05))
        tf_bullets = bullets_box.text_frame
        tf_bullets.word_wrap = True
        tf_bullets.vertical_anchor = MSO_ANCHOR.MIDDLE

        bullets = data.get("bullets", [])
        total_lines = len(bullets)
        base_pt = 13.5 if total_lines <= 4 else (12.5 if total_lines <= 6 else 11.5)
        space_after = Pt(10) if total_lines <= 4 else (Pt(6) if total_lines <= 6 else Pt(4))

        for b_idx, bullet in enumerate(bullets):
            p = tf_bullets.paragraphs[0] if b_idx == 0 else tf_bullets.add_paragraph()
            p.space_after = space_after
            p.line_spacing = 1.15
            
            is_sub = bullet.strip().startsWith('•') or bullet.strip().startsWith('-') or bullet.strip().startsWith('–') or bullet.startswith('  ')
            clean_b = bullet.lstrip(' •-–0123456789.)')
            
            if is_sub:
                p.level = 1
                p.text = f"– {clean_b}"
                p.font.size = Pt(base_pt - 1.5)
                p.font.color.rgb = SUBTEXT_COLOR
            else:
                p.level = 0
                p.text = f"• {clean_b}"
                p.font.size = Pt(base_pt)
                p.font.color.rgb = TEXT_COLOR
                
            p.font.name = "微軟正黑體"

        # Right side Image Card
        img_card_x, img_card_y = Inches(7.55), Inches(0.9)
        img_card_w, img_card_h = Inches(5.18), Inches(5.9)
        img_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_card_x, img_card_y, img_card_w, img_card_h)
        set_shape_flat_color(img_card, CARD_BG, BORDER_COLOR)

        img_rel = data.get("image", "")
        img_full = os.path.join(base_assets_dir, img_rel) if not os.path.isabs(img_rel) else img_rel

        if os.path.exists(img_full):
            try:
                slide.shapes.add_picture(img_full, img_card_x + Inches(0.08), img_card_y + Inches(0.08), img_card_w - Inches(0.16), img_card_h - Inches(0.16))
            except Exception as e:
                print(f"[WARN] Failed to insert image {img_full}: {e}")

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12.133), Inches(0.35))
        tf_f = footer_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = f"頁碼 {s_idx + 1} / {len(slides_data)}  |  學習共同體課例探究"
        p_f.font.size = Pt(10)
        p_f.font.color.rgb = SUBTEXT_COLOR
        p_f.font.name = "微軟正黑體"

        # Speaker notes
        if "notes" in data and data["notes"]:
            notes_slide = slide.notes_slide
            tf_n = notes_slide.notes_text_frame
            tf_n.text = data["notes"]

    prs.save(output_pptx_path)
    print(f"[SUCCESS] PPTX Deck generated at: {output_pptx_path}")


def export_pdf_from_html(html_path, output_pdf_path):
    edge_paths = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"
    ]
    msedge = next((p for p in edge_paths if os.path.exists(p)), None)
    if not msedge:
        print("[WARN] Microsoft Edge not found, skipping direct HTML-to-PDF export.")
        return

    cmd = [
        msedge,
        "--headless",
        "--disable-gpu",
        f"--print-to-pdf={output_pdf_path}",
        "--no-pdf-header-footer",
        html_path
    ]
    try:
        subprocess.run(cmd, check=True)
        print(f"[SUCCESS] PDF exported at: {output_pdf_path}")
    except Exception as e:
        print(f"[WARN] PDF export failed: {e}")


def main():
    parser = argparse.ArgumentParser(description="Build Responsive HTML, PPTX and PDF Lesson Study Decks")
    parser.add_argument("--config", required=True, help="Path to case study JSON data")
    parser.add_argument("--output-dir", default=".", help="Directory to save generated files")
    parser.add_argument("--prefix", default="lesson_study", help="Prefix for output filenames")
    args = parser.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        slides_data = json.load(f)

    os.makedirs(args.output_dir, exist_ok=True)
    html_path = os.path.join(args.output_dir, f"{args.prefix}.html")
    pptx_path = os.path.join(args.output_dir, f"{args.prefix}.pptx")
    pdf_path = os.path.join(args.output_dir, f"{args.prefix}.pdf")

    build_html_deck(slides_data, html_path)
    build_pptx_deck(slides_data, pptx_path, base_assets_dir=os.path.dirname(os.path.abspath(args.config)))
    export_pdf_from_html(html_path, pdf_path)

if __name__ == "__main__":
    main()
